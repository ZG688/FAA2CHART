# -*- coding: utf-8 -*-
"""
Analysis of three-task ablation experiment results & statistics table export
"""
from __future__ import annotations

import os
import sys
import json
import math
from collections import defaultdict
from typing import Any, Dict, List, Tuple, Optional

# ── shared statistical utilities ──────────────────────────────────
_HERE = os.path.dirname(os.path.abspath(__file__))
_EVAL_ROOT = os.path.normpath(os.path.join(_HERE, "..", ".."))  # evaluation/
if _EVAL_ROOT not in sys.path:
    sys.path.insert(0, _EVAL_ROOT)
from stat_utils import (          # noqa: E402
    fmt_mean_sd, fmt_median_iqr, cliffs_delta, mann_whitney_p, fmt_p,
)

_ROOT = os.environ.get("ABLATION_ROOT", "data")
_NEW_OUT = os.environ.get("FAA2CHART_ABLATION_OUT", os.path.join(_ROOT, "eval_results"))

TASKS_CN = {"simple": "Simple tasks", "moderate": "Moderate tasks", "complex": "Complex tasks"}
VARIANTS_CN = {"multi": "Multi-agent", "single": "Single-agent"}

METRICS = [
    ("elapsed_sec",        "Elapsed (s)",        "mean_sd",    True),   # lower is better
    ("total_steps",        "Total steps",        "mean_sd",    False),
    ("tool_calls",         "Tool calls",         "mean_sd",    True),   # fewer is better
    ("subagent_delegations","Sub-agent delegations","mean_sd", False),
    ("tool_errors",        "Tool errors",        "mean_sd",    True),
    ("observable_points",  "Observable points",  "mean_sd",    False),
    ("recovery_rate",      "Recovery rate",      "mean",       False),
    ("call_result_pairing_rate","Pairing rate",  "mean",       False),
]


# ==================================================================
# Statistics utilities (shared functions imported from stat_utils)
# ==================================================================
def _p_interpretation(p: Optional[float], multi_vals: List[float],
                      single_vals: List[float], label: str,
                      low_better: bool) -> str:
    """Translate P value into textual interpretation."""
    if p is None or (isinstance(p, float) and math.isnan(p)) or p >= 0.05:
        return "No significant difference"
    m_med = sorted(multi_vals)[len(multi_vals) // 2] if multi_vals else 0
    s_med = sorted(single_vals)[len(single_vals) // 2] if single_vals else 0
    if abs(m_med - s_med) < 1e-9:
        return "No significant difference"

    # Determine which agent is better
    if low_better:
        better = "Multi" if m_med < s_med else "Single"
    else:
        better = "Multi" if m_med > s_med else "Single"

    # Determine direction word: for the better agent, is the metric higher or lower?
    if label in ("Elapsed (s)",):
        dir_word = "faster"
    elif label in ("Tool errors", "Tool calls"):
        dir_word = "fewer"
    elif low_better:
        dir_word = "fewer"
    else:
        dir_word = "more"

    return f"{better}-agent {dir_word}"


# ==================================================================
# Data loading
# ==================================================================
_DIR_ALIAS = {
    "simple":  ["simple", "simple_tasks"],
    "moderate":["moderate", "moderate_tasks"],
    "complex": ["complex", "complex_tasks"],
}
_VARIANT_ALIAS = {
    "multi":  ["multi", "multi_agent"],
    "single": ["single", "single_agent"],
}


def _resolve_dir(base: str, candidates: List[str]) -> str | None:
    for c in candidates:
        p = os.path.join(base, c)
        if os.path.isdir(p):
            return p
    return None


def _read_all() -> Dict[str, Dict[str, List[Dict]]]:
    data: Dict[str, Dict[str, List[Dict]]] = defaultdict(lambda: defaultdict(list))
    for task in ["simple", "moderate", "complex"]:
        task_dir = _resolve_dir(_NEW_OUT, _DIR_ALIAS[task])
        if task_dir is None:
            continue
        for agent in ["multi", "single"]:
            ag_dir = _resolve_dir(task_dir, _VARIANT_ALIAS[agent])
            if ag_dir is None:
                continue
            for fn in sorted(os.listdir(ag_dir)):
                if not fn.startswith("run_") or not fn.endswith(".json"):
                    continue
                with open(os.path.join(ag_dir, fn), "r", encoding="utf-8") as f:
                    rec = json.load(f)
                data[task][agent].append(rec)
    return data


def _flatten_metric(rec: Dict, metric_key: str) -> float | None:
    if metric_key == "elapsed_sec":
        v = rec.get("elapsed_sec", -1)
        return None if (v is None or v < 0) else float(v)
    m = rec.get("metrics", {})
    v = m.get(metric_key)
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ==================================================================
# Table generation
# ==================================================================
def _task_summary_table(task: str, data: Dict[str, List[Dict]]) -> str:
    lines: List[str] = []
    cn = TASKS_CN[task]
    lines.append(f"### {task} ({cn}) — Single/Multi-agent comparison table")
    lines.append("")
    header = "| Metric | Multi-agent | Single-agent | P-value | P-value interpretation |"
    lines.append(header)
    lines.append("|---|---|---|---|---|")
    for key, label, mode, low_better in METRICS:
        multi_vals = [_flatten_metric(r, key) for r in data["multi"] if r.get("success")]
        single_vals = [_flatten_metric(r, key) for r in data["single"] if r.get("success")]
        multi_vals = [v for v in multi_vals if v is not None]
        single_vals = [v for v in single_vals if v is not None]

        if mode == "median_iqr":
            multi_s = fmt_median_iqr(multi_vals)
            single_s = fmt_median_iqr(single_vals)
        elif mode == "mean":
            m_m = (sum(multi_vals) / len(multi_vals)) if multi_vals else None
            s_m = (sum(single_vals) / len(single_vals)) if single_vals else None
            multi_s = f"{m_m:.3f}" if m_m is not None else "n/a"
            single_s = f"{s_m:.3f}" if s_m is not None else "n/a"
        else:  # mean_sd
            multi_s = fmt_mean_sd(multi_vals)
            single_s = fmt_mean_sd(single_vals)

        p = mann_whitney_p(multi_vals, single_vals) if multi_vals and single_vals else None
        interp = _p_interpretation(p, multi_vals, single_vals, label, low_better)
        lines.append(f"| {label} | {multi_s} | {single_s} | {fmt_p(p)} | {interp} |")

    # Success rate
    m_total = len(data["multi"])
    m_success = sum(1 for r in data["multi"] if r.get("success"))
    s_total = len(data["single"])
    s_success = sum(1 for r in data["single"] if r.get("success"))
    m_rate = f"{m_success}/{m_total} ({100*m_success/m_total:.0f}%)" if m_total else "n/a"
    s_rate = f"{s_success}/{s_total} ({100*s_success/s_total:.0f}%)" if s_total else "n/a"
    try:
        from scipy.stats import fisher_exact
        table = [[m_success, m_total - m_success],
                 [s_success, s_total - s_success]]
        _, fp = fisher_exact(table)
        fp_s = fmt_p(fp)
    except Exception:
        fp_s = "n/a"
    lines.append(f"| Success rate | {m_rate} | {s_rate} | {fp_s} | — |")
    lines.append("")
    lines.append("*Note: P-values from two-sided Mann-Whitney U test (Fisher's exact test for success rate); "
                 "\\*, \\*\\*, \\*\\*\\* denote P<0.05, P<0.01, P<0.001 respectively; "
                 "n.s. indicates not significant. *")
    lines.append("")
    return "\n".join(lines)


def _overview_table(data) -> str:
    lines = ["### Overview: Three task types × Single/Multi-agent success rate & time comparison table", ""]
    lines.append("| Task | Variant | n (success) | Success rate | Elapsed (mean±SD) | Tool calls (mean±SD) |")
    lines.append("|---|---|---|---|---|---|")
    for task in ["simple", "moderate", "complex"]:
        for agent in ["multi", "single"]:
            rs = data[task][agent]
            if not rs:
                continue
            succ = [r for r in rs if r.get("success")]
            total = len(rs)
            rate = f"{len(succ)}/{total} ({100*len(succ)/total:.0f}%)" if total else "n/a"
            el = [_flatten_metric(r, "elapsed_sec") for r in succ]
            el = [v for v in el if v is not None]
            tc = [_flatten_metric(r, "tool_calls") for r in succ]
            tc = [v for v in tc if v is not None]
            lines.append(f"| {TASKS_CN[task]} | {VARIANTS_CN[agent]} | {total} ({len(succ)}) | {rate} | "
                         f"{fmt_mean_sd(el)} | {fmt_mean_sd(tc)} |")
    lines.append("")
    return "\n".join(lines)


# ==================================================================
# docx export
# ==================================================================
def _export_docx(text: str):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        import re
    except Exception as e:
        print(f"[WARNING] Cannot generate docx: {e}")
        return

    doc = Document()
    for run in doc.styles['Normal'].font.runs if False else []:
        pass
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)

    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=2)
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
            continue
        if line.startswith('|') and any('---' in l for l in lines[i:i+1] if l):
            # Read markdown table until non-| line
            tbl = [line]
            while i < len(lines) and lines[i].rstrip().startswith('|'):
                tbl.append(lines[i].rstrip())
                i += 1
            # Filter separator lines
            data_lines = [l for l in tbl if not re.match(r'^\|[\s\-:|]+\|$', l.strip())]
            if len(data_lines) < 2:
                continue
            headers = [c.strip() for c in data_lines[0].split('|')[1:-1]]
            rows_data = [[c.strip() for c in l.split('|')[1:-1]] for l in data_lines[1:]]
            t = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(headers):
                cell = t.rows[0].cells[j]
                cell.text = ''
                run = cell.paragraphs[0].add_run(h)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9)
                run.bold = True
            for ri, row in enumerate(rows_data):
                for j, c in enumerate(row):
                    if j < len(headers):
                        cell = t.rows[1 + ri].cells[j]
                        cell.text = ''
                        run = cell.paragraphs[0].add_run(c)
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(9)
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)

    out_path = os.path.join(_NEW_OUT, "ablation_stats_report.docx")
    doc.save(out_path)
    print(f"[docx generated] {out_path}")


def main():
    data = _read_all()
    parts: List[str] = []
    parts.append("# Three-task ablation experiment — Statistical analysis report")
    parts.append("")
    parts.append(_overview_table(data))
    for task in ["simple", "moderate", "complex"]:
        if data[task]["multi"] or data[task]["single"]:
            parts.append(_task_summary_table(task, data[task]))

    txt = "\n".join(parts)
    out_txt = os.path.join(_NEW_OUT, "ablation_stats_report.txt")
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write(txt)
    print(f"[txt generated] {out_txt}")
    print(txt)
    _export_docx(txt)


if __name__ == "__main__":
    main()